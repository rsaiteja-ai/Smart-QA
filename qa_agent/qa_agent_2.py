#!/usr/bin/env python3
"""Prompt-pack driven QA test case generator.

Inputs:
- Excel requirement workbook
- DOCX requirement document
- Image/screenshot

Generation:
- Loads rules from QA_Agent_Prompt_Pack/*.md and output_schema.json.
- Calls an OpenAI-compatible LLM Foundry endpoint when QA_AGENT_URL and
  QA_AGENT_API_KEY are configured.
- Falls back to a deterministic generator so the app still works offline.
"""

import argparse
import base64
import json
import os
import re
import urllib.error
import urllib.request
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


BASE_DIR = Path(__file__).resolve().parent
PROMPT_PACK_DIR = BASE_DIR / "QA_Agent_Prompt_Pack"
DEFAULT_MODEL = "gpt-4o-mini"

OUTPUT_COLUMNS = [
    "Pre-condition",
    "Feature Mapping",
    "Test Case ID",
    "Test Case Description",
    "Summary Description",
    "Test Step ID",
    "Test Steps",
    "Test Data",
    "Expected Result",
    "Actual Result",
]


def clean_text(value):
    if value is None or pd.isna(value):
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return re.sub(r"[ \t]+", " ", text).strip()


def normalize_key(value):
    return re.sub(r"[^a-z0-9]+", "", str(value).lower())


def find_column(df, candidates):
    normalized = {normalize_key(column): column for column in df.columns}
    for candidate in candidates:
        key = normalize_key(candidate)
        if key in normalized:
            return normalized[key]
    for candidate in candidates:
        key = normalize_key(candidate)
        for norm, column in normalized.items():
            if key in norm:
                return column
    return None


def split_items(value):
    text = clean_text(value)
    if not text:
        return []
    items = []
    for line in text.split("\n"):
        item = re.sub(r"^\s*[-*•\d.)]+", "", line).strip()
        if item:
            items.append(item)
    return items or [text]


def read_text_file(path):
    return path.read_text(encoding="utf-8").strip() if path.exists() else ""


def load_prompt_pack():
    parts = []
    for name in [
        "system_prompt.md",
        "testcase_rules.md",
        "excel_rules.md",
        "naming_rules.md",
        "review_rules.md",
        "examples.md",
    ]:
        path = PROMPT_PACK_DIR / name
        if path.exists():
            parts.append(f"\n\n--- {name} ---\n{read_text_file(path)}")
    schema_path = PROMPT_PACK_DIR / "output_schema.json"
    if schema_path.exists():
        parts.append(f"\n\n--- output_schema.json ---\n{read_text_file(schema_path)}")
    return "\n".join(parts).strip()


def load_config(override=None):
    config = {}
    override = override or {}
    config_path = BASE_DIR / "qa_agent_config.json"
    if config_path.exists():
        try:
            config = json.loads(config_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            config = {}
    return {
        "url": override.get("url") or os.getenv("QA_AGENT_URL") or config.get("LLM_URL") or config.get("url"),
        "api_key": (
            override.get("api_key")
            or os.getenv("QA_AGENT_API_KEY")
            or config.get("API_KEY")
            or config.get("api_key")
        ),
        "model": override.get("model") or os.getenv("QA_AGENT_MODEL") or config.get("MODEL") or config.get("model") or DEFAULT_MODEL,
    }


def requirements_from_excel(excel_path):
    workbook = pd.ExcelFile(excel_path, engine="openpyxl")
    sheet_name = workbook.sheet_names[0]
    df = pd.read_excel(workbook, sheet_name=sheet_name)
    records = []
    for index, row in df.iterrows():
        values = {
            str(column): clean_text(row.get(column, ""))
            for column in df.columns
            if clean_text(row.get(column, ""))
        }
        if values:
            records.append({"row": int(index) + 1, "values": values})
    return json.dumps({"source_type": "excel", "requirements": records}, indent=2)


def requirements_from_docx(docx_path):
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(docx_path) as docx:
        xml = docx.read("word/document.xml")
    root = ET.fromstring(xml)
    paragraphs = []
    for paragraph in root.findall(".//w:p", namespace):
        text = "".join(node.text or "" for node in paragraph.findall(".//w:t", namespace)).strip()
        if text:
            paragraphs.append(text)
    tables = []
    for table in root.findall(".//w:tbl", namespace):
        table_rows = []
        for row in table.findall(".//w:tr", namespace):
            cells = []
            for cell in row.findall(".//w:tc", namespace):
                cell_text = " ".join(
                    node.text or "" for node in cell.findall(".//w:t", namespace)
                ).strip()
                cells.append(cell_text)
            if cells:
                table_rows.append(cells)
        if table_rows:
            tables.append(table_rows)
    return json.dumps({"source_type": "docx", "paragraphs": paragraphs, "tables": tables}, indent=2)


def requirements_from_image(image_path):
    suffix = Path(image_path).suffix.lower().lstrip(".") or "png"
    data = Path(image_path).read_bytes()
    encoded = base64.b64encode(data).decode("ascii")
    return json.dumps({
        "source_type": "image",
        "filename": Path(image_path).name,
        "image_base64": encoded,
        "instruction": "Infer requirements from this Figma/application screenshot.",
    })


def build_generation_prompt(requirements_text):
    prompt_pack = load_prompt_pack()
    return f"""
{prompt_pack}

You must generate manual functional test cases from the requirements below.

Return only valid JSON, without markdown fences. The JSON must be:
{{
  "test_cases": [
    {{
      "pre_condition": "",
      "feature_mapping": "",
      "test_case_id": "",
      "test_case_description": "",
      "summary_description": "",
      "steps": [
        {{
          "test_step_id": "",
          "test_steps": "",
          "test_data": "",
          "expected_result": "",
          "actual_result": ""
        }}
      ]
    }}
  ]
}}

Rules:
- Follow every rule in QA_Agent_Prompt_Pack.
- Repeat common fields only on the first row of each test case when converted to Excel.
- Use the naming pattern from naming_rules.md.
- Actual Result must be blank.
- Produce positive, negative, boundary, and validation scenarios.
- Use professional QA language.

Requirements:
{requirements_text}
""".strip()


def normalize_llm_url(url):
    if not url:
        return ""
    if url.rstrip("/").endswith("/chat/completions"):
        return url.rstrip("/")
    return f"{url.rstrip('/')}/chat/completions"


def call_llm(prompt, llm_config=None):
    config = load_config(llm_config)
    if not config["url"] or not config["api_key"]:
        return None

    payload = {
        "model": config["model"],
        "messages": [
            {"role": "system", "content": "You are a Senior QA Engineer. Return only valid JSON."},
            {"role": "user", "content": prompt},
        ],
        "temperature": 0.2,
    }
    request = urllib.request.Request(
        normalize_llm_url(config["url"]),
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {config['api_key']}",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=90) as response:
            raw = response.read().decode("utf-8")
    except urllib.error.URLError as exc:
        raise RuntimeError(f"LLM Foundry request failed: {exc}") from exc

    body = json.loads(raw)
    content = body["choices"][0]["message"]["content"]
    return parse_llm_json(content)


def test_llm_connection(llm_config=None):
    config = load_config(llm_config)
    if not config["url"] or not config["api_key"]:
        raise ValueError("LLM endpoint and API key are required.")

    payload = {
        "model": config["model"],
        "messages": [
            {"role": "system", "content": "Return only valid JSON."},
            {"role": "user", "content": "Return {\"ok\": true} if you can read this message."},
        ],
        "temperature": 0,
    }
    request = urllib.request.Request(
        normalize_llm_url(config["url"]),
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {config['api_key']}",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            body = json.loads(response.read().decode("utf-8"))
    except urllib.error.URLError as exc:
        raise RuntimeError(f"LLM Foundry connection failed: {exc}") from exc

    content = body["choices"][0]["message"]["content"].strip()
    if content.startswith("```"):
        content = re.sub(r"^```(?:json)?", "", content, flags=re.IGNORECASE).strip()
        content = re.sub(r"```$", "", content).strip()
    json.loads(content)
    return {"status": "ok", "model": config["model"]}


def parse_llm_json(content):
    text = content.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?", "", text, flags=re.IGNORECASE).strip()
        text = re.sub(r"```$", "", text).strip()
    match = re.search(r"\{.*\}", text, flags=re.DOTALL)
    if match:
        text = match.group(0)
    data = json.loads(text)
    if isinstance(data, list):
        data = {"test_cases": data}
    if "test_cases" not in data:
        raise ValueError("LLM response did not include a test_cases array.")
    return data


def fallback_cases(requirements_text):
    data = json.loads(requirements_text)
    records = data.get("requirements") or [{"values": {"Feature Name": "Requirement"}}]
    test_cases = []
    for index, record in enumerate(records, start=1):
        values = record.get("values", {})
        feature = (
            values.get("Feature Name")
            or values.get("Feature")
            or values.get("Module")
            or f"Feature {index}"
        )
        page = re.sub(r"[^A-Za-z0-9]+", "", feature.title()) or f"Feature{index}"
        user_story = values.get("User Story", "")
        acceptance = split_items(values.get("Acceptance criteria") or values.get("Acceptance Criteria") or "")
        ui = split_items(values.get("UI Elements", ""))
        rules = split_items(values.get("Business Rules", ""))
        data_text = values.get("Test Data", "")
        scenarios = [
            ("positive", f"Verify {feature} works with valid data", "Valid data must be accepted."),
            ("negative", f"Validate {feature} rejects invalid data", "Invalid data must be rejected with a clear message."),
            ("boundary", f"Verify {feature} handles boundary values", "Boundary values must be handled according to business rules."),
            ("validation", f"Confirm mandatory validations for {feature}", "Required field validation must be displayed."),
        ]
        for case_index, (_, description, expected) in enumerate(scenarios, start=1):
            case_id = f"{page}-TC_{case_index:03d}"
            test_cases.append({
                "pre_condition": f"User has access to {feature}.",
                "feature_mapping": feature,
                "test_case_id": case_id,
                "test_case_description": description,
                "summary_description": f"{description}. {user_story}".strip(),
                "steps": [
                    {
                        "test_step_id": f"{case_id}-01",
                        "test_steps": f"Navigate to {feature}.",
                        "test_data": "",
                        "expected_result": f"{feature} page or workflow must be available.",
                        "actual_result": "",
                    },
                    {
                        "test_step_id": f"{case_id}-02",
                        "test_steps": first_non_empty([
                            f"Use {ui[0]}." if ui else "",
                            f"Apply rule: {rules[0]}." if rules else "",
                            acceptance[0] if acceptance else "",
                            description,
                        ]),
                        "test_data": data_text,
                        "expected_result": expected,
                        "actual_result": "",
                    },
                ],
            })
    return {"test_cases": test_cases}


def first_non_empty(values):
    for value in values:
        if value:
            return value
    return ""


def flatten_cases(data):
    rows = []
    for test_case in data.get("test_cases", []):
        steps = test_case.get("steps") or []
        if not steps:
            steps = [{}]
        for index, step in enumerate(steps):
            first = index == 0
            rows.append({
                "Pre-condition": clean_text(test_case.get("pre_condition", "")) if first else "",
                "Feature Mapping": clean_text(test_case.get("feature_mapping", "")) if first else "",
                "Test Case ID": clean_text(test_case.get("test_case_id", "")) if first else "",
                "Test Case Description": clean_text(test_case.get("test_case_description", "")) if first else "",
                "Summary Description": clean_text(test_case.get("summary_description", "")) if first else "",
                "Test Step ID": clean_text(step.get("test_step_id", "")),
                "Test Steps": clean_text(step.get("test_steps", "")),
                "Test Data": clean_text(step.get("test_data", "")),
                "Expected Result": clean_text(step.get("expected_result", "")),
                "Actual Result": "",
            })
    return rows


def write_to_excel(data, output_path):
    rows = flatten_cases(data)
    if not rows:
        raise ValueError("No test cases were generated.")
    output = Path(output_path)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(rows, columns=OUTPUT_COLUMNS).to_excel(writer, sheet_name="Sheet", index=False)
    format_workbook(output)
    return output


def format_workbook(path):
    workbook = load_workbook(path)
    worksheet = workbook.active
    header_fill = PatternFill("solid", fgColor="1D4ED8")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = {
        "A": 32, "B": 24, "C": 22, "D": 42, "E": 46,
        "F": 22, "G": 54, "H": 34, "I": 54, "J": 20,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    worksheet.freeze_panes = "A2"
    workbook.save(path)


def generate_from_requirements(requirements_text, output_path, llm_config=None):
    prompt = build_generation_prompt(requirements_text)
    generated = call_llm(prompt, llm_config)
    if generated is None:
        generated = fallback_cases(requirements_text)
    return write_to_excel(generated, output_path)


def generate_from_excel(excel_path, output_path=None, llm_config=None):
    input_path = Path(excel_path)
    output = Path(output_path) if output_path else input_path.with_name(f"{input_path.stem}_generated_test_cases.xlsx")
    return generate_from_requirements(requirements_from_excel(input_path), output, llm_config)


def generate_from_docx(docx_path, output_path=None, llm_config=None):
    input_path = Path(docx_path)
    output = Path(output_path) if output_path else input_path.with_name(f"{input_path.stem}_generated_test_cases.xlsx")
    return generate_from_requirements(requirements_from_docx(input_path), output, llm_config)


def generate_from_image(image_path, output_path=None, llm_config=None):
    input_path = Path(image_path)
    output = Path(output_path) if output_path else input_path.with_name(f"{input_path.stem}_generated_test_cases.xlsx")
    return generate_from_requirements(requirements_from_image(input_path), output, llm_config)


def main():
    parser = argparse.ArgumentParser(description="Smart QA Test Case Generator")
    parser.add_argument("--excel", help="Path to Excel requirement file")
    parser.add_argument("--docx", help="Path to DOCX requirement file")
    parser.add_argument("--image", help="Path to image/screenshot requirement file")
    parser.add_argument("--output", default=None, help="Path for generated output Excel file")
    parser.add_argument("--no-interactive", action="store_true", help="Run without interactive prompts")
    args = parser.parse_args()

    selected = [value for value in [args.excel, args.docx, args.image] if value]
    if len(selected) != 1:
        raise SystemExit("ERROR: Provide exactly one of --excel, --docx, or --image.")
    if args.excel:
        output = generate_from_excel(args.excel, args.output)
    elif args.docx:
        output = generate_from_docx(args.docx, args.output)
    else:
        output = generate_from_image(args.image, args.output)
    print(f"Test cases generated successfully: {output}")


if __name__ == "__main__":
    main()
