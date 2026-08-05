#!/usr/bin/env python3
"""
Simple polling watcher for AI_QA_Test_Case_Generator.xlsx.
When any cell value or new row is added, runs qa_agent_2.py.
Usage: python watch_excel_trigger.py
Requires: openpyxl (pip install openpyxl)
"""
import hashlib
import os
import subprocess
import sys
import time

try:
    from openpyxl import load_workbook
except Exception:
    print("openpyxl is required. Install with: pip install openpyxl")
    sys.exit(1)

FILE = "AI_QA_Test_Case_Generator.xlsx"
POLL_INTERVAL = 5  # seconds


def workbook_hash(path):
    # Build deterministic text of all sheets/cells and hash it
    wb = load_workbook(path, data_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f"--SHEET:{sheet_name}--")
        for row in ws.iter_rows(values_only=True):
            # represent row cells separated by |; None -> empty string
            parts.append("|".join("" if v is None else str(v) for v in row))
    text = "\n".join(parts)
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def try_hash(path, attempts=6, delay=0.5):
    # retry a few times to handle file being saved/locked
    last_exc = None
    for _ in range(attempts):
        try:
            return workbook_hash(path)
        except Exception as e:
            last_exc = e
            time.sleep(delay)
    raise last_exc


def trigger_agent(script_path):
    # Run qa_agent_2.py with same Python interpreter
    print(f"Triggering {script_path}")
    try:
        proc = subprocess.run([sys.executable, script_path], capture_output=True, text=True)
        print(f"Exit code: {proc.returncode}")
        if proc.stdout:
            print("---stdout---\n", proc.stdout)
        if proc.stderr:
            print("---stderr---\n", proc.stderr)
    except Exception as e:
        print("Failed to run agent:", e)


def main():
    base = os.path.dirname(os.path.abspath(__file__))
    path = os.path.join(base, FILE)
    agent = os.path.join(base, "qa_agent_2.py")

    if not os.path.exists(path):
        print(f"Excel file not found: {path}")
        sys.exit(1)
    if not os.path.exists(agent):
        print(f"Agent script not found: {agent}")
        sys.exit(1)

    try:
        last_hash = try_hash(path)
    except Exception as e:
        print("Unable to read workbook on startup:", e)
        sys.exit(1)

    print(f"Watching {path} (poll every {POLL_INTERVAL}s). Ctrl-C to stop.")

    try:
        while True:
            time.sleep(POLL_INTERVAL)
            try:
                cur_hash = try_hash(path)
            except Exception as e:
                print("Read error, will retry:", e)
                continue
            if cur_hash != last_hash:
                print("Change detected in workbook content.")
                trigger_agent(agent)
                last_hash = cur_hash
    except KeyboardInterrupt:
        print("Watcher stopped by user.")


if __name__ == "__main__":
    main()
