#!/usr/bin/env python3
"""
Event-driven watcher for AI_QA_Test_Case_Generator.xlsx using watchdog.
Triggers qa_agent_2.py when any cell value or new row is added/edited.
Usage: pip install watchdog openpyxl
       python watch_excel_watchdog.py
"""
import hashlib
import os
import subprocess
import sys
import threading
import time

try:
    from openpyxl import load_workbook
except Exception:
    print("openpyxl is required. Install with: pip install openpyxl")
    sys.exit(1)

try:
    from watchdog.observers import Observer
    from watchdog.events import FileSystemEventHandler
except Exception:
    print("watchdog is required. Install with: pip install watchdog")
    sys.exit(1)

FILE = "AI_QA_Test_Case_Generator.xlsx"
DEBOUNCE = 1.0  # seconds to wait for burst of events to settle


def workbook_hash(path):
    wb = load_workbook(path, data_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f"--SHEET:{sheet_name}--")
        for row in ws.iter_rows(values_only=True):
            parts.append("|".join("" if v is None else str(v) for v in row))
    text = "\n".join(parts)
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def try_hash(path, attempts=8, delay=0.5):
    last_exc = None
    for _ in range(attempts):
        try:
            return workbook_hash(path)
        except Exception as e:
            last_exc = e
            time.sleep(delay)
    raise last_exc


def trigger_agent(script_path, excel_path):
    print(f"Triggering {script_path} for {excel_path}")
    try:
        proc = subprocess.run(
            [sys.executable, script_path, "--excel", excel_path, "--no-interactive"],
            capture_output=True,
            text=True,
            encoding="utf-8",
            env={**os.environ, "PYTHONUTF8": "1"},
            cwd=os.path.dirname(script_path)
        )
        print(f"Exit code: {proc.returncode}")
        if proc.stdout:
            print("---stdout---\n" + proc.stdout)
        if proc.stderr:
            print("---stderr---\n" + proc.stderr)
    except Exception as e:
        print("Failed to run agent:", e)


class ExcelChangeHandler(FileSystemEventHandler):
    def __init__(self, target_path, agent_path):
        super().__init__()
        self.target_path = os.path.normcase(os.path.abspath(target_path))
        self.agent_path = agent_path
        self._timer = None
        self._lock = threading.Lock()
        self.last_hash = None
        self._agent_thread = None
        # initialize hash if file exists
        if os.path.exists(self.target_path):
            try:
                self.last_hash = try_hash(self.target_path)
            except Exception as e:
                print("Initial read failed:", e)

    def _schedule_check(self):
        with self._lock:
            if self._timer:
                self._timer.cancel()
            self._timer = threading.Timer(DEBOUNCE, self._on_debounced_event)
            self._timer.daemon = True
            self._timer.start()

    def on_modified(self, event):
        if os.path.normcase(os.path.abspath(event.src_path)) == self.target_path:
            print("Modified event received for target file")
            self._schedule_check()

    def on_created(self, event):
        if os.path.normcase(os.path.abspath(event.src_path)) == self.target_path:
            print("Created event received for target file")
            self._schedule_check()

    def _on_debounced_event(self):
        try:
            cur_hash = try_hash(self.target_path)
        except Exception as e:
            print("Read error after event, will ignore this burst:", e)
            return
        if cur_hash != self.last_hash:
            print("Change detected in workbook content (event-driven).")
            if self._agent_thread and self._agent_thread.is_alive():
                print("Agent is already running; skipping this trigger.")
            else:
                self._agent_thread = threading.Thread(
                    target=trigger_agent,
                    args=(self.agent_path, self.target_path),
                    daemon=True
                )
                self._agent_thread.start()
            self.last_hash = cur_hash
        else:
            print("No content change detected after event.")


def main():
    base = os.path.dirname(os.path.abspath(__file__))
    path = os.path.join(base, FILE)
    agent = os.path.join(base, "qa_agent_2.py")

    if not os.path.exists(agent):
        print(f"Agent script not found: {agent}")
        sys.exit(1)

    directory = os.path.dirname(path) or "."

    event_handler = ExcelChangeHandler(path, agent)
    observer = Observer()
    observer.schedule(event_handler, directory, recursive=False)
    observer.start()

    print(f"Watching {path} for changes (event-driven). Ctrl-C to stop.")
    try:
        while True:
            time.sleep(1)
    except KeyboardInterrupt:
        print("Stopping observer...")
        observer.stop()
    observer.join()


if __name__ == "__main__":
    main()
